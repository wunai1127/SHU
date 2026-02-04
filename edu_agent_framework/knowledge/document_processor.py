"""
文档处理器
==========
处理教师上传的各类教学文档
"""

from typing import List, Dict, Any, Optional
from pathlib import Path
from dataclasses import dataclass
import json


@dataclass
class DocumentChunk:
    """文档块"""
    content: str
    source: str
    page: int = 0
    chunk_index: int = 0
    metadata: Dict[str, Any] = None

    def to_dict(self) -> dict:
        return {
            "content": self.content,
            "source": self.source,
            "page": self.page,
            "chunk_index": self.chunk_index,
            "metadata": self.metadata or {}
        }


class DocumentProcessor:
    """
    文档处理器

    支持格式：
    - PDF: 教材、讲义
    - PPTX: PPT课件
    - DOCX: Word文档
    - TXT/MD: 纯文本/Markdown
    - JSON/YAML: 结构化数据

    处理流程：
    1. 解析文档 → 2. 分块 → 3. 清洗 → 4. 返回结构化数据
    """

    def __init__(
        self,
        chunk_size: int = 500,
        chunk_overlap: int = 50,
        language: str = "chinese"
    ):
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap
        self.language = language

    def process(self, file_path: str) -> List[DocumentChunk]:
        """
        处理单个文档

        Args:
            file_path: 文档路径

        Returns:
            文档块列表
        """
        path = Path(file_path)
        suffix = path.suffix.lower()

        processors = {
            '.pdf': self._process_pdf,
            '.pptx': self._process_pptx,
            '.ppt': self._process_pptx,
            '.docx': self._process_docx,
            '.doc': self._process_docx,
            '.txt': self._process_text,
            '.md': self._process_text,
            '.json': self._process_json,
            '.yaml': self._process_yaml,
            '.yml': self._process_yaml,
        }

        processor = processors.get(suffix)
        if not processor:
            raise ValueError(f"不支持的文件格式: {suffix}")

        return processor(str(path))

    def process_batch(self, file_paths: List[str]) -> List[DocumentChunk]:
        """批量处理文档"""
        all_chunks = []
        for path in file_paths:
            try:
                chunks = self.process(path)
                all_chunks.extend(chunks)
            except Exception as e:
                print(f"处理文件失败 {path}: {e}")
        return all_chunks

    def _process_pdf(self, file_path: str) -> List[DocumentChunk]:
        """处理PDF文档"""
        try:
            from PyPDF2 import PdfReader
        except ImportError:
            raise ImportError("请安装PyPDF2: pip install PyPDF2")

        reader = PdfReader(file_path)
        chunks = []

        for page_num, page in enumerate(reader.pages):
            text = page.extract_text()
            if text:
                page_chunks = self._split_text(text)
                for i, chunk in enumerate(page_chunks):
                    chunks.append(DocumentChunk(
                        content=chunk,
                        source=file_path,
                        page=page_num + 1,
                        chunk_index=i,
                        metadata={"type": "pdf"}
                    ))

        return chunks

    def _process_pptx(self, file_path: str) -> List[DocumentChunk]:
        """处理PPT文档"""
        try:
            from pptx import Presentation
        except ImportError:
            raise ImportError("请安装python-pptx: pip install python-pptx")

        prs = Presentation(file_path)
        chunks = []

        for slide_num, slide in enumerate(prs.slides):
            slide_text = []
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    slide_text.append(shape.text)

            if slide_text:
                text = "\n".join(slide_text)
                slide_chunks = self._split_text(text)
                for i, chunk in enumerate(slide_chunks):
                    chunks.append(DocumentChunk(
                        content=chunk,
                        source=file_path,
                        page=slide_num + 1,
                        chunk_index=i,
                        metadata={"type": "pptx", "slide": slide_num + 1}
                    ))

        return chunks

    def _process_docx(self, file_path: str) -> List[DocumentChunk]:
        """处理Word文档"""
        try:
            from docx import Document
        except ImportError:
            raise ImportError("请安装python-docx: pip install python-docx")

        doc = Document(file_path)
        text = "\n".join([para.text for para in doc.paragraphs])

        chunks = []
        text_chunks = self._split_text(text)
        for i, chunk in enumerate(text_chunks):
            chunks.append(DocumentChunk(
                content=chunk,
                source=file_path,
                page=0,
                chunk_index=i,
                metadata={"type": "docx"}
            ))

        return chunks

    def _process_text(self, file_path: str) -> List[DocumentChunk]:
        """处理纯文本文档"""
        with open(file_path, 'r', encoding='utf-8') as f:
            text = f.read()

        chunks = []
        text_chunks = self._split_text(text)
        for i, chunk in enumerate(text_chunks):
            chunks.append(DocumentChunk(
                content=chunk,
                source=file_path,
                page=0,
                chunk_index=i,
                metadata={"type": "text"}
            ))

        return chunks

    def _process_json(self, file_path: str) -> List[DocumentChunk]:
        """处理JSON结构化数据"""
        with open(file_path, 'r', encoding='utf-8') as f:
            data = json.load(f)

        chunks = []
        if isinstance(data, list):
            for i, item in enumerate(data):
                content = json.dumps(item, ensure_ascii=False)
                chunks.append(DocumentChunk(
                    content=content,
                    source=file_path,
                    page=0,
                    chunk_index=i,
                    metadata={"type": "json", "item_index": i}
                ))
        else:
            content = json.dumps(data, ensure_ascii=False)
            text_chunks = self._split_text(content)
            for i, chunk in enumerate(text_chunks):
                chunks.append(DocumentChunk(
                    content=chunk,
                    source=file_path,
                    page=0,
                    chunk_index=i,
                    metadata={"type": "json"}
                ))

        return chunks

    def _process_yaml(self, file_path: str) -> List[DocumentChunk]:
        """处理YAML结构化数据"""
        import yaml
        with open(file_path, 'r', encoding='utf-8') as f:
            data = yaml.safe_load(f)

        # 转换为JSON处理
        content = json.dumps(data, ensure_ascii=False)
        chunks = []
        text_chunks = self._split_text(content)
        for i, chunk in enumerate(text_chunks):
            chunks.append(DocumentChunk(
                content=chunk,
                source=file_path,
                page=0,
                chunk_index=i,
                metadata={"type": "yaml"}
            ))

        return chunks

    def _split_text(self, text: str) -> List[str]:
        """
        文本分块

        使用简单的滑动窗口方法，支持中文
        """
        if not text:
            return []

        # 清理文本
        text = text.strip()

        # 如果文本较短，直接返回
        if len(text) <= self.chunk_size:
            return [text]

        chunks = []
        start = 0

        while start < len(text):
            end = start + self.chunk_size

            # 尝试在句子边界切分
            if end < len(text):
                # 寻找最近的句子结束符
                for sep in ['。', '！', '？', '\n', '.', '!', '?']:
                    pos = text.rfind(sep, start, end)
                    if pos > start:
                        end = pos + 1
                        break

            chunk = text[start:end].strip()
            if chunk:
                chunks.append(chunk)

            start = end - self.chunk_overlap

        return chunks


class FAQExtractor:
    """
    FAQ提取器

    从教学文档中自动提取问答对
    """

    def __init__(self, llm_client=None):
        self.llm_client = llm_client

    def extract_from_text(self, text: str, num_qa: int = 5) -> List[Dict[str, str]]:
        """
        使用LLM从文本中提取FAQ

        Args:
            text: 输入文本
            num_qa: 要提取的问答对数量

        Returns:
            FAQ列表
        """
        if not self.llm_client:
            return []

        prompt = f"""请从以下教学内容中提取{num_qa}个学生可能会问的问题及其答案。

教学内容：
{text[:2000]}

请以JSON格式返回，每个问答对包含"question"和"answer"字段。
"""

        try:
            response = self.llm_client.chat.completions.create(
                model="gpt-4o",
                messages=[{"role": "user", "content": prompt}],
                response_format={"type": "json_object"}
            )
            result = json.loads(response.choices[0].message.content)
            return result.get("faqs", result.get("questions", []))
        except Exception as e:
            print(f"FAQ提取失败: {e}")
            return []


class ConceptExtractor:
    """
    概念提取器

    从教学文档中提取核心概念
    """

    def __init__(self, llm_client=None):
        self.llm_client = llm_client

    def extract_concepts(self, text: str) -> List[str]:
        """从文本中提取核心概念"""
        if not self.llm_client:
            # 简单的关键词提取
            return self._simple_extract(text)

        prompt = f"""请从以下教学内容中提取核心概念/术语（不超过20个）：

{text[:3000]}

请以JSON格式返回，包含"concepts"字段（字符串列表）。
"""

        try:
            response = self.llm_client.chat.completions.create(
                model="gpt-4o",
                messages=[{"role": "user", "content": prompt}],
                response_format={"type": "json_object"}
            )
            result = json.loads(response.choices[0].message.content)
            return result.get("concepts", [])
        except Exception as e:
            print(f"概念提取失败: {e}")
            return self._simple_extract(text)

    def _simple_extract(self, text: str) -> List[str]:
        """简单的关键词提取（不依赖LLM）"""
        try:
            import jieba.analyse
            keywords = jieba.analyse.extract_tags(text, topK=20)
            return keywords
        except ImportError:
            return []
